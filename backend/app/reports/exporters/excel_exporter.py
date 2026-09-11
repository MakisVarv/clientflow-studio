# type: ignore
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font


class ExcelExporter:

    @staticmethod
    def export(sheet_name, data):

        workbook = Workbook()

        worksheet = workbook.active

        worksheet.title = sheet_name

        # Header
        worksheet["A1"] = "Field"
        worksheet["B1"] = "Value"

        worksheet["A1"].font = Font(bold=True)
        worksheet["B1"].font = Font(bold=True)

        row = 2

        for key, value in data.items():

            worksheet.cell(row=row, column=1).value = key

            worksheet.cell(row=row, column=2).value = value

            row += 1

        buffer = BytesIO()

        workbook.save(buffer)

        buffer.seek(0)

        return buffer
